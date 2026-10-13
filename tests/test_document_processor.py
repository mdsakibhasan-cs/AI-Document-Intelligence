from pathlib import Path

from docx import Document

from src.document_processor import extract_text


def test_extract_text_from_pdf():
    pdf_path = Path("data/sample.pdf")

    text = extract_text(pdf_path)

    assert isinstance(text, str)
    assert len(text) > 100
    assert "AIM" in text


def test_extract_text_from_docx(tmp_path):
    docx_path = tmp_path / "test_document.docx"

    document = Document()
    document.add_paragraph(
        "This is a test document for document extraction."
    )
    document.add_paragraph(
        "It contains multiple paragraphs."
    )
    document.save(docx_path)

    text = extract_text(docx_path)

    assert isinstance(text, str)
    assert "This is a test document" in text
    assert "multiple paragraphs" in text


def test_unsupported_file_type():
    unsupported_file = Path("data/test.txt")

    try:
        extract_text(unsupported_file)
        assert False
    except ValueError as error:
        assert "Only PDF and DOCX files are supported." in str(error)